#!/usr/bin/env python3
"""CLI: run many-to-one exact-cent matching on a two-tab (GL/Bank-style)
workbook and write results back.

Writes:
  - group code into the source sheets' `--id-col` column (both GL and
    Bank rows in a matched group get the same code), and the engine tag
    into `--engine-col`.
  - a "Match Report" tab (one row per target: status, matched sum,
    difference, item count, source row refs, components) and a
    "Match Inputs" tab (one row per source item: assigned target, match
    report row) modeled on the Excel Matching Base example workbook.

Usage:
  python3 match_workbook.py --input SAP_202512_Unmatched.xlsx \\
      --output SAP_202512_Matched.xlsx --prefix SAP --period 2025-12
"""

import argparse
import copy
from datetime import date, datetime

import openpyxl
from openpyxl.styles import Font

from subset_match import match_all, to_cents, MatchResult

STATUS_LABELS = {
    "exact_1to1": "Exact match (1:1)",
    "exact_many": "Exact match (many-to-one)",
    "not_allocated": "Not allocated",
    "search_incomplete": "Search incomplete (pool too large)",
}


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise ValueError(f"Unrecognized date value: {value!r}")


def load_side(ws, amount_col: str, date_col: str, id_prefix: str):
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    if amount_col not in idx:
        raise SystemExit(f"Column '{amount_col}' not found on sheet '{ws.title}'")
    items = []
    for row_num, row in enumerate(
        ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True), start=2
    ):
        amt = row[idx[amount_col]]
        if amt is None:
            continue
        dt = row[idx[date_col]]
        items.append({
            "id": f"{id_prefix}{row_num}",
            "row": row_num,
            "amount_cents": to_cents(amt),
            "date": _as_date(dt),
        })
    return items, idx


def write_back(ws, idx, id_col: str, engine_col: str, engine_tag: str,
                assignments: dict, status_col: str = None, status_value: str = None):
    """assignments: {row_num: group_code}"""
    id_col_idx = idx[id_col] + 1  # openpyxl is 1-indexed
    engine_col_idx = idx[engine_col] + 1
    status_col_idx = idx[status_col] + 1 if status_col else None
    for row_num, code in assignments.items():
        ws.cell(row=row_num, column=id_col_idx, value=code)
        ws.cell(row=row_num, column=engine_col_idx, value=engine_tag)
        if status_col_idx:
            ws.cell(row=row_num, column=status_col_idx, value=status_value)


def format_components(amounts: list) -> str:
    parts = []
    for i, a in enumerate(amounts):
        val = f"${abs(a):,.2f}"
        if i == 0:
            parts.append(f"-{val}" if a < 0 else val)
        else:
            parts.append(f"{'-' if a < 0 else '+'} {val}")
    return " ".join(parts)


def build_report_tabs(wb, gl_items, bank_items, results: dict, group_codes: dict):
    if "Match Report" in wb.sheetnames:
        del wb["Match Report"]
    if "Match Inputs" in wb.sheetnames:
        del wb["Match Inputs"]
    rep = wb.create_sheet("Match Report")
    inputs = wb.create_sheet("Match Inputs")

    bold = Font(bold=True)
    rep.append(["Many-to-One Exact Match Report"])
    rep.append(["Global exact-cent allocation within the configured date window; "
                 "each GL row used at most once across all targets; "
                 "maximum items per target enforced by the engine run."])
    rep.append([])
    header = ["Target Row", "Target Amount", "Status", "Matched Sum",
              "Difference", "Item Count", "Group Code", "GL Rows", "GL Amounts"]
    rep.append(header)
    for c in rep[4]:
        c.font = bold

    gl_by_id = {item["id"]: item for item in gl_items}
    bank_by_row = {item["row"]: item for item in bank_items}

    report_row_for_bank_row = {}
    for b in sorted(bank_items, key=lambda x: x["row"]):
        r: MatchResult = results[b["id"]]
        target_dollars = b["amount_cents"] / 100
        status_label = STATUS_LABELS[r.status]
        if r.group_ids:
            gl_rows = [gl_by_id[gid]["row"] for gid in r.group_ids]
            gl_amounts = [gl_by_id[gid]["amount_cents"] / 100 for gid in r.group_ids]
            matched_sum = sum(gl_amounts)
            diff = round(matched_sum - target_dollars, 2)
            count = len(gl_rows)
            code = group_codes.get(b["row"], "")
            components = format_components(gl_amounts)
            rep.append([b["row"], target_dollars, status_label, matched_sum, diff,
                        count, code, ", ".join(str(r) for r in gl_rows), components])
        else:
            rep.append([b["row"], target_dollars, status_label, None, None,
                        None, "", "", ""])
        report_row_for_bank_row[b["row"]] = rep.max_row

    inputs.append(["Many-to-One Match Inputs"])
    inputs.append(["Source (GL) rows with their assigned target and match report row."])
    inputs.append([])
    ih = ["GL Row", "Amount", "Group Code", "Assigned Target Amount", "Match Report Row"]
    inputs.append(ih)
    for c in inputs[4]:
        c.font = bold

    row_to_target = {}
    for b in bank_items:
        r: MatchResult = results[b["id"]]
        for gid in r.group_ids:
            row_to_target[gl_by_id[gid]["row"]] = b

    for item in sorted(gl_items, key=lambda x: x["row"]):
        target = row_to_target.get(item["row"])
        if target is not None:
            code = group_codes.get(target["row"], "")
            inputs.append([item["row"], item["amount_cents"] / 100, code,
                            target["amount_cents"] / 100,
                            report_row_for_bank_row[target["row"]]])
        else:
            inputs.append([item["row"], item["amount_cents"] / 100, "", "Unused", None])


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--gl-sheet", default="GL")
    ap.add_argument("--bank-sheet", default="Bank")
    ap.add_argument("--amount-col", default="Matching Amount")
    ap.add_argument("--date-col", default="Transaction Date")
    ap.add_argument("--id-col", default="Match ID or Group ID")
    ap.add_argument("--engine-col", default="Matched By Engine")
    ap.add_argument("--engine-tag", default="Manual")
    ap.add_argument("--status-col", default=None, help="e.g. 'Match Status'")
    ap.add_argument("--status-value", default="Matched")
    ap.add_argument("--prefix", required=True, help="e.g. SAP")
    ap.add_argument("--period", required=True, help="e.g. 2025-12")
    ap.add_argument("--max-items", type=int, default=10)
    ap.add_argument("--date-window", type=int, default=15)
    ap.add_argument("--cap-entries", type=int, default=2_000_000,
                     help="DP entry budget per target search before giving up")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.input, data_only=False)
    gl_ws = wb[args.gl_sheet]
    bank_ws = wb[args.bank_sheet]

    gl_items, gl_idx = load_side(gl_ws, args.amount_col, args.date_col, "GL")
    bank_items, bank_idx = load_side(bank_ws, args.amount_col, args.date_col, "BK")

    print(f"Loaded {len(gl_items)} GL rows, {len(bank_items)} Bank rows")

    results = match_all(gl_items, bank_items, max_items=args.max_items,
                         date_window_days=args.date_window, cap_entries=args.cap_entries)

    matched_bank_rows = sorted(
        (b["row"] for b in bank_items if results[b["id"]].group_ids),
    )
    group_codes = {
        row: f"{args.prefix}-{args.period}-{seq:03d}"
        for seq, row in enumerate(matched_bank_rows, start=1)
    }

    gl_by_id = {item["id"]: item for item in gl_items}
    gl_assignments = {}
    bank_assignments = {}
    for b in bank_items:
        r = results[b["id"]]
        if not r.group_ids:
            continue
        code = group_codes[b["row"]]
        bank_assignments[b["row"]] = code
        for gid in r.group_ids:
            gl_assignments[gl_by_id[gid]["row"]] = code

    write_back(gl_ws, gl_idx, args.id_col, args.engine_col, args.engine_tag, gl_assignments,
               args.status_col, args.status_value)
    write_back(bank_ws, bank_idx, args.id_col, args.engine_col, args.engine_tag, bank_assignments,
               args.status_col, args.status_value)

    build_report_tabs(wb, gl_items, bank_items, results, group_codes)

    exact = sum(1 for r in results.values() if r.group_ids)
    print(f"Matched {exact} of {len(bank_items)} Bank targets "
          f"({sum(1 for r in results.values() if r.status == 'exact_1to1')} 1:1, "
          f"{sum(1 for r in results.values() if r.status == 'exact_many')} many-to-one)")

    wb.save(args.output)
    print(f"Saved: {args.output}")


if __name__ == "__main__":
    main()
