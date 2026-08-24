#!/usr/bin/env python3
"""Add an 'Analysis' dashboard tab (KPI band, narrative, two charts, a
transaction-mix table) to a workbook that has a Bank-style sheet, modeled
on the Excel Matching Base example's Analysis tab.

Usage:
  python3 build_analysis_tab.py --input file.xlsx --output file.xlsx \\
      --bank-sheet Bank
"""

import argparse
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "122447"
BLUE = "4A76B8"
LIGHT_BLUE = "EAF1FA"
WHITE = "FFFFFF"


def fmt_money(v):
    sign = "-" if v < 0 else ""
    return f"{sign}${abs(v):,.2f}"


def build(input_path, output_path, bank_sheet, currency, source_label):
    wb = openpyxl.load_workbook(input_path, data_only=False)
    bank = wb[bank_sheet]
    idx = {c.value: i for i, c in enumerate(next(bank.iter_rows(min_row=1, max_row=1)))}
    last_row = bank.max_row
    for req in ("Transaction Date", "Transaction Type", "Debit", "Credit",
                "Matching Amount", "Description"):
        if req not in idx:
            raise SystemExit(f"Bank sheet missing required column: {req}")

    date_col = get_column_letter(idx["Transaction Date"] + 1)
    type_col = get_column_letter(idx["Transaction Type"] + 1)
    debit_col = get_column_letter(idx["Debit"] + 1)
    credit_col = get_column_letter(idx["Credit"] + 1)
    amt_col = get_column_letter(idx["Matching Amount"] + 1)
    date_range = f"{bank_sheet}!${date_col}$2:${date_col}${last_row}"
    type_range = f"{bank_sheet}!${type_col}$2:${type_col}${last_row}"
    debit_range = f"{bank_sheet}!${debit_col}$2:${debit_col}${last_row}"
    credit_range = f"{bank_sheet}!${credit_col}$2:${credit_col}${last_row}"
    amt_range = f"{bank_sheet}!${amt_col}$2:${amt_col}${last_row}"

    rows = list(bank.iter_rows(min_row=2, max_row=last_row, values_only=True))
    rows = [r for r in rows if r[idx["Transaction Date"]] is not None]
    n_txn = len(rows)
    dates = sorted({r[idx["Transaction Date"]].date() for r in rows})
    net_total = sum(r[idx["Matching Amount"]] for r in rows)
    debit_total = -sum(r[idx["Debit"]] for r in rows if r[idx["Debit"]] is not None)
    credit_total = sum(r[idx["Credit"]] for r in rows if r[idx["Credit"]] is not None)

    by_type = defaultdict(float)
    type_count = defaultdict(int)
    for r in rows:
        t = r[idx["Transaction Type"]] or "(blank)"
        by_type[t] += r[idx["Matching Amount"]]
        type_count[t] += 1
    types_sorted = sorted(by_type.items(), key=lambda x: -abs(x[1]))

    gross = sum(abs(r[idx["Matching Amount"]]) for r in rows)
    top5 = sorted(rows, key=lambda r: -abs(r[idx["Matching Amount"]]))[:5]
    top5_pct = sum(abs(r[idx["Matching Amount"]]) for r in top5) / gross * 100 if gross else 0
    top_movers = sorted(rows, key=lambda r: -abs(r[idx["Matching Amount"]]))[:len(types_sorted)]

    status_col = idx.get("Match Status")
    if status_col is not None:
        unmatched = sum(1 for r in rows if r[status_col] == "Unmatched")
        matched = n_txn - unmatched
    else:
        unmatched = matched = None

    lead_types = types_sorted[:3]
    lead_desc = ", ".join(
        f"{t.title()} {'net outflow' if v < 0 else 'net inflow'} of {fmt_money(abs(v))}"
        for t, v in lead_types
    )
    def fmt_date(d):
        return f"{d.strftime('%b')} {d.day}, {d.year}"

    date_span = f"{fmt_date(dates[0])} - {fmt_date(dates[-1])}" if dates else ""
    story = (
        f"This population covers {n_txn} Bank transactions ({date_span}) with a "
        f"{fmt_money(abs(net_total))} net {'outflow' if net_total < 0 else 'inflow'}: "
        f"{fmt_money(debit_total)} of debits {'exceeded' if debit_total > credit_total else 'were below'} "
        f"{fmt_money(credit_total)} of credits. Activity is concentrated -- the five largest "
        f"transactions represent {top5_pct:.1f}% of gross movement. {lead_desc}."
    )
    if unmatched is not None:
        story += (
            f" Of the {n_txn} transactions, {matched} have been matched by the many-to-one "
            f"engine so far; {unmatched} remain Unmatched and need manual or rule-engine review."
        )

    if "Analysis" in wb.sheetnames:
        del wb["Analysis"]
    ws = wb.create_sheet("Analysis", 0)

    for col, width in zip("ABCDEFGHI", (14, 12, 14, 3, 14, 16, 10, 8, 14)):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:I1")
    ws["A1"] = f"{source_label} Bank Activity Analysis"
    ws["A1"].font = Font(size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 34
    for col in "ABCDEFGHI":
        ws[f"{col}1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center", horizontal="left", indent=1)

    ws.merge_cells("A2:I2")
    ws["A2"] = f"{n_txn} exported transactions | {date_span} | {currency} | Source: {bank_sheet}"
    ws["A2"].font = Font(italic=True, size=10)
    ws.row_dimensions[2].height = 18

    kpis = [
        ("A", "TRANSACTIONS", n_txn, "#,##0"),
        ("C", "NET CASH FLOW", net_total, '"$"#,##0'),
        ("E", "TOTAL DEBITS", debit_total, '"$"#,##0'),
        ("G", "TOTAL CREDITS", credit_total, '"$"#,##0'),
    ]
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 26
    ws.row_dimensions[6].height = 16
    for col, label, value, numfmt in kpis:
        end_col = chr(ord(col) + 1)
        ws.merge_cells(f"{col}4:{end_col}4")
        c = ws[f"{col}4"]
        c.value = label
        c.font = Font(bold=True, size=9, color=NAVY)
        c.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.merge_cells(f"{col}5:{end_col}6")
        v = ws[f"{col}5"]
        v.value = value
        v.number_format = numfmt
        v.font = Font(bold=True, size=18, color=NAVY)
        v.alignment = Alignment(horizontal="center", vertical="center")
        v.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        for r in (4, 5, 6):
            ws[f"{end_col}{r}"].fill = PatternFill("solid", fgColor=LIGHT_BLUE)

    ws.merge_cells("A8:I8")
    ws["A8"] = "STORY IN BRIEF"
    ws["A8"].font = Font(bold=True, color=WHITE, size=11)
    ws["A8"].fill = PatternFill("solid", fgColor=BLUE)
    for col in "ABCDEFGHI":
        ws[f"{col}8"].fill = PatternFill("solid", fgColor=BLUE)

    ws.merge_cells("A9:I12")
    ws["A9"] = story
    ws["A9"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    ws["A9"].font = Font(size=10)

    # --- Helper columns for the daily chart: J Date, K Net, L Credits, M Debits ---
    ws["J1"], ws["K1"], ws["L1"], ws["M1"] = "Date", "Net Flow", "Credits", "Debits"
    for i, d in enumerate(dates):
        r = 2 + i
        ws[f"J{r}"] = datetime(d.year, d.month, d.day)
        ws[f"J{r}"].number_format = "mm/dd"
        ws[f"K{r}"] = (f"=SUMIFS({bank_sheet}!${amt_col}$2:${amt_col}${last_row},"
                        f"{bank_sheet}!${date_col}$2:${date_col}${last_row},J{r})")
        ws[f"L{r}"] = (f"=SUMIFS({bank_sheet}!${credit_col}$2:${credit_col}${last_row},"
                        f"{bank_sheet}!${date_col}$2:${date_col}${last_row},J{r})")
        ws[f"M{r}"] = (f"=-SUMIFS({bank_sheet}!${debit_col}$2:${debit_col}${last_row},"
                        f"{bank_sheet}!${date_col}$2:${date_col}${last_row},J{r})")

    line = LineChart()
    line.title = "Daily Net Cash Flow"
    line.height, line.width = 8, 17
    line.y_axis.title = "Net cash flow (USD)"
    line.x_axis.title = "Posting date"
    line.x_axis.number_format = "mm/dd"
    line.x_axis.axPos = "b"
    line.y_axis.axPos = "l"
    line.x_axis.delete = False
    line.y_axis.delete = False
    data = Reference(ws, min_col=11, min_row=1, max_row=1 + len(dates))
    cats = Reference(ws, min_col=10, min_row=2, max_row=1 + len(dates))
    line.add_data(data, titles_from_data=True)
    line.set_categories(cats)
    line.series[0].smooth = False
    ws.add_chart(line, "A14")

    # --- Transaction mix table (also feeds the bar chart) ---
    mix_start = 50
    ws.merge_cells(f"A{mix_start-2}:I{mix_start-2}")
    ws[f"A{mix_start-2}"] = "TRANSACTION MIX AND LARGEST MOVEMENTS"
    ws[f"A{mix_start-2}"].font = Font(bold=True, color=WHITE, size=11)
    ws[f"A{mix_start-2}"].fill = PatternFill("solid", fgColor=BLUE)
    for col in "ABCDEFGHI":
        ws[f"{col}{mix_start-2}"].fill = PatternFill("solid", fgColor=BLUE)

    headers = ["Type", "Transactions", "Net Amount", "", "Post Date", "Description", "Type", "D/C", "Amount"]
    for c, h in zip("ABCDEFGHI", headers):
        cell = ws[f"{c}{mix_start-1}"]
        cell.value = h
        cell.font = Font(bold=True)

    for i, (t, v) in enumerate(types_sorted):
        r = mix_start + i
        ws[f"A{r}"] = t
        ws[f"B{r}"] = type_count[t]
        ws[f"C{r}"] = round(v, 2)
        ws[f"C{r}"].number_format = '"$"#,##0.00;("$"#,##0.00)'

    for i, m in enumerate(top_movers):
        r = mix_start + i
        ws[f"E{r}"] = m[idx["Transaction Date"]]
        ws[f"E{r}"].number_format = "yyyy-mm-dd"
        ws[f"F{r}"] = (m[idx["Description"]] or "")[:60]
        ws[f"G{r}"] = m[idx["Transaction Type"]]
        ws[f"H{r}"] = "Credit" if m[idx["Matching Amount"]] >= 0 else "Debit"
        ws[f"I{r}"] = round(m[idx["Matching Amount"]], 2)
        ws[f"I{r}"].number_format = '"$"#,##0.00;("$"#,##0.00)'

    bar = BarChart()
    bar.type = "bar"
    bar.title = "Net Cash Flow by Transaction Type"
    bar.height, bar.width = 8, 17
    bar.x_axis.title = "Transaction type"
    bar.y_axis.title = "Net cash flow (USD)"
    bar.x_axis.axPos = "l"  # bar.type="bar" draws horizontal bars: category axis is vertical
    bar.y_axis.axPos = "b"
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bdata = Reference(ws, min_col=3, min_row=mix_start - 1, max_row=mix_start - 1 + len(types_sorted))
    bcats = Reference(ws, min_col=1, min_row=mix_start, max_row=mix_start - 1 + len(types_sorted))
    bar.add_data(bdata, titles_from_data=True)
    bar.set_categories(bcats)
    ws.add_chart(bar, "A31")

    wb.save(output_path)
    print(f"Saved: {output_path}")
    print(f"Transactions={n_txn} Net={net_total:.2f} Debits={debit_total:.2f} Credits={credit_total:.2f}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True)
    ap.add_argument("--output", required=True)
    ap.add_argument("--bank-sheet", default="Bank")
    ap.add_argument("--currency", default="USD")
    ap.add_argument("--source-label", default="SAP")
    args = ap.parse_args()
    build(args.input, args.output, args.bank_sheet, args.currency, args.source_label)


if __name__ == "__main__":
    main()
